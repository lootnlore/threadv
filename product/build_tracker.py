"""
Builds the paid Reseller Tracker workbook (named by tracker.name in site.config.mjs).

    python3 product/build_tracker.py   -> product/dist/ThreadVet-Reseller-Tracker.xlsx

If LibreOffice (`soffice`) is installed, the file is recalculated so every
formula also carries a stored value (spreadsheet previews show numbers, not
blanks). Excel and Google Sheets recalculate on open either way.
product/verify_tracker.py checks the fee formulas against the website engine.

The marketplace list, fee rates, default sales tax and IRS mileage rates are
read from the website's source (src/engine/fees.mjs, src/engine/calc.mjs,
src/data/mileage.mjs) through Node, so a change there is picked up on the next
build. The output file is sold, so product/dist/ is git-ignored: never commit
the .xlsx to this public repository.

Formulas stick to Excel 2010-era functions (SUMIFS, COUNTIFS, AVERAGEIFS,
INDEX/MATCH, CHOOSE, IFERROR) so the workbook behaves the same in Excel 2016+,
Google Sheets and LibreOffice.
"""

import json
import re
import shutil
import subprocess
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = Path(__file__).parent / "dist"
DATA_ROWS = 1500  # pre-built rows per log; copy the last row down for more
FIRST = 5  # first data row on every log sheet
LAST = FIRST + DATA_ROWS - 1
RANGE_END = 5000  # summaries read rows FIRST..RANGE_END, so logs can grow past the prebuilt rows

# ---------- styles ----------
FONT = "Arial"
BRAND = "0B6B58"
INK = "16201D"
MUTED = "56615D"
INPUT_FILL = PatternFill("solid", fgColor="FFF7DB")
CALC_FILL = PatternFill("solid", fgColor="F1F3F2")
HEAD_FILL = PatternFill("solid", fgColor=BRAND)
SOFT_FILL = PatternFill("solid", fgColor="E3F3EE")
THIN = Side(style="thin", color="D9D6CF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

f_title = Font(name=FONT, size=18, bold=True, color=INK)
f_sub = Font(name=FONT, size=10, color=MUTED)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_label = Font(name=FONT, size=10, bold=True, color=INK)
f_body = Font(name=FONT, size=10, color=INK)
f_input = Font(name=FONT, size=10, color="0000FF")
f_calc = Font(name=FONT, size=10, color="000000")
f_kpi = Font(name=FONT, size=16, bold=True, color=BRAND)
f_section = Font(name=FONT, size=12, bold=True, color=BRAND)
f_note = Font(name=FONT, size=9, italic=True, color=MUTED)

USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
PCT = '0%;(0%);"-"'
DATE_FMT = "mm/dd/yyyy"
INT = '#,##0;(#,##0);"-"'

# The marketplaces come from the website engine, in its order, plus this
# spreadsheet-only one for local cash sales (no fees unless the user sets some).
OTHER = "Other / Local"

SOURCES = [
    "Goodwill",
    "Salvation Army",
    "Other thrift store",
    "Garage & yard sales",
    "Estate sales",
    "Facebook Marketplace",
    "Retail arbitrage",
    "Online arbitrage",
    "Other",
]

EXPENSE_CATEGORIES = [
    "Shipping supplies",
    "Postage (not on sale labels)",
    "Software & subscriptions",
    "Marketplace ads & store plans",
    "Sourcing trips (non-mileage)",
    "Equipment",
    "Office supplies",
    "Phone & internet (business share)",
    "Education",
    "Other",
]

# (name in formulas, platform id, setting, value, format, note). A string value is
# a key of that platform's RATES in src/engine/fees.mjs, so a fee change is made
# once, there. OTHER's rows (platform None) exist only in the spreadsheet and keep
# a literal default. Notes may use {handbags}, {media}, {boost} and {other}.
FEES = [
    ("eBay_FVF", "ebay", "Final value fee (most categories)", "fvf", "0.00%", "On item + shipping + sales tax. Handbags {handbags}, books/media {media}: use Actual fees for those."),
    ("eBay_FVF_Cap", "ebay", "Final value fee applies up to", "fvfCap", USD0, "Order total above this is charged the rate below."),
    ("eBay_FVF_Over", "ebay", "Rate above that amount", "fvfOver", "0.00%", ""),
    ("eBay_OrderFee", "ebay", "Per-order fee", "orderFee", USD, "Orders over the small-order limit."),
    ("eBay_OrderFeeSmall", "ebay", "Per-order fee, small orders", "orderFeeSmall", USD, ""),
    ("eBay_SmallOrderMax", "ebay", "Small-order limit (order total)", "smallOrderMax", USD, ""),
    ("Posh_Rate", "poshmark", "Commission at or above threshold", "rate", "0.00%", "Buyer pays the shipping label."),
    ("Posh_Flat", "poshmark", "Flat fee below threshold", "flat", USD, ""),
    ("Posh_Threshold", "poshmark", "Threshold (sale price)", "threshold", USD, ""),
    ("Mercari_Rate", "mercari", "Selling fee", "rate", "0.00%", "On item + shipping. No seller processing fee since Jan 2025."),
    ("Depop_Proc", "depop", "Payment processing", "proc", "0.00%", "US sellers: 0% selling fee. On item + shipping + tax."),
    ("Depop_ProcFixed", "depop", "Processing fixed fee", "procFixed", USD, "Boosted sales add {boost}: log it under Actual fees."),
    ("Etsy_Listing", "etsy", "Listing fee", "listing", USD, ""),
    ("Etsy_Txn", "etsy", "Transaction fee", "txn", "0.00%", "On item + shipping."),
    ("Etsy_Proc", "etsy", "Payment processing", "proc", "0.00%", "On order total incl. tax."),
    ("Etsy_ProcFixed", "etsy", "Processing fixed fee", "procFixed", USD, "Offsite Ads sales: log under Actual fees."),
    ("Whatnot_Comm", "whatnot", "Commission (on item price)", "commission", "0.00%", "Lower at $15k+ per 4 weeks since Sep 21, 2026."),
    ("Whatnot_Proc", "whatnot", "Payment processing", "proc", "0.00%", "On item + shipping + tax."),
    ("Whatnot_ProcFixed", "whatnot", "Processing fixed fee", "procFixed", USD, ""),
    ("FB_Rate", "facebook", "Selling fee (shipped orders)", "rate", "0.00%", "On item + shipping. Local pickup: choose {other}."),
    ("FB_Min", "facebook", "Minimum fee", "min", USD, ""),
    ("Grailed_Threshold", "grailed", "Lower-rate threshold (sale price)", "threshold", USD, "Rate cut on May 20, 2026."),
    ("Grailed_LowRate", "grailed", "Commission below threshold", "lowRate", "0.00%", ""),
    ("Grailed_Min", "grailed", "Minimum commission below threshold", "min", USD, ""),
    ("Grailed_Rate", "grailed", "Commission at or above threshold", "rate", "0.00%", ""),
    ("Grailed_Proc", "grailed", "Payment processing (US)", "proc", "0.00%", "On item + shipping."),
    ("Grailed_ProcFixed", "grailed", "Processing fixed fee", "procFixed", USD, ""),
    ("TikTok_Rate", "tiktok", "Referral fee", "referral", "0.00%", "On item + buyer-paid shipping, before tax. Add creator commission if any."),
    ("Other_Rate", None, "Fee rate", 0.0, "0.00%", "Local cash sales: leave at 0."),
    ("Other_Fixed", None, "Fixed fee per sale", 0.0, USD, ""),
]

def engine_info():
    """Marketplaces, fee rates, default sales tax and mileage rates, straight from the website's source."""
    script = (
        "import { PLATFORMS, FEES_VERIFIED, RATES, pctText } from './src/engine/fees.mjs';"
        "import { DEFAULTS } from './src/engine/calc.mjs';"
        "import { IRS_MILEAGE_RATES, IRS_MILEAGE_SOURCE } from './src/data/mileage.mjs';"
        "import config from './site.config.mjs';"
        "console.log(JSON.stringify({ brand: config.name, product: config.tracker.name,"
        " verified: FEES_VERIFIED, rates: RATES, taxRate: DEFAULTS.taxRate,"
        " mileage: IRS_MILEAGE_RATES, mileageSource: IRS_MILEAGE_SOURCE,"
        " platforms: PLATFORMS.map((p) => ({ id: p.id, name: p.name, sellerPaysShipping: p.sellerPaysShipping,"
        " feesIncludeTax: Boolean(p.feesIncludeTax), source: p.sources[0].url })),"
        # Percentages quoted in Fees tab notes, formatted exactly as the website does.
        " text: { handbags: pctText(RATES.ebay.handbagsFvf), media: pctText(RATES.ebay.mediaFvf),"
        " boost: pctText(RATES.depop.boost) } }));"
    )
    try:
        out = subprocess.run(["node", "--input-type=module", "-e", script], cwd=ROOT, check=True, capture_output=True, text=True)
    except FileNotFoundError:
        raise SystemExit("Node.js 22 or newer is needed: the tracker reads its fees from the website engine.") from None
    except subprocess.CalledProcessError as err:
        raise SystemExit(f"Could not read the website engine (src/engine, src/data):\n{err.stderr.strip()}") from None
    info = json.loads(out.stdout)
    day = datetime.strptime(info["verified"], "%Y-%m-%d")
    info["verified_label"] = f"{day:%B} {day.day}, {day.year}"  # no %-d: it fails on Windows
    return info


ENGINE = {}


def engine():
    """engine_info(), loaded once."""
    if not ENGINE:
        ENGINE.update(engine_info())
    return ENGINE


def platform_names():
    """The Platforms list: every website marketplace in engine order, then OTHER."""
    return [p["name"] for p in engine()["platforms"]] + [OTHER]


def and_list(items):
    """['eBay', 'Etsy', 'Depop'] -> 'eBay, Etsy and Depop' (same style as andList on the website)."""
    return items[0] if len(items) == 1 else f"{', '.join(items[:-1])} and {items[-1]}" if items else ""


def as_of(days_before=0):
    """A date relative to the fee verification date, so example rows stay recent."""
    return date.fromisoformat(engine()["verified"]) - timedelta(days=days_before)


def buyer_ships():
    """1-based Platforms positions where the buyer pays the label (shipping inputs ignored)."""
    return [i + 1 for i, p in enumerate(engine()["platforms"]) if not p["sellerPaysShipping"]]


def buyer_ships_test(cell):
    """Formula that is TRUE when `cell` holds a marketplace where the buyer pays the label."""
    tests = [f"{cell}=INDEX(Platforms,{i})" for i in buyer_ships()]
    return "FALSE" if not tests else tests[0] if len(tests) == 1 else f"OR({','.join(tests)})"


def buyer_ships_names():
    return and_list([p["name"] for p in engine()["platforms"] if not p["sellerPaysShipping"]])


def tax_fee_names():
    return and_list([p["name"] for p in engine()["platforms"] if p["feesIncludeTax"]])


# Prices chosen to hit every fee cliff: eBay's $10 small order and $7,500 cap,
# Poshmark's $15, Grailed's $120 and $1.99 minimum, Facebook's $0.80 minimum,
# and half-cent rounding cases.
VERIFY_PRICES = [5, 9.25, 12.99, 14.99, 15, 19.99, 45, 45.05, 99.95, 119.99, 120, 333.33, 7499.99, 8000]
VERIFY_SHIPPING = [0, 4.5, 7.99]
VERIFY_COST = 3
VERIFY_LABEL = 6.5


NO_FEE_CASES = [None, "Bogus Market"]  # Sold on left blank, and a name not in Platforms
NO_DATE_CASES = ["Mercari"]  # a sale entered without its date sold
NO_PRICE_CASES = ["Depop"]  # a date sold entered without the sale price


def no_fee_row(i):
    return FIRST + 1 + len(verify_cases()) + i


def no_date_row(i):
    return no_fee_row(len(NO_FEE_CASES)) + i


def no_price_row(i):
    return no_date_row(len(NO_DATE_CASES)) + i


def verify_cases():
    return [(m, p, s) for m in platform_names() for p in VERIFY_PRICES for s in VERIFY_SHIPPING]


def style(cell, font=f_body, fill=None, fmt=None, align=None, border=None):
    cell.font = font
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    return cell


def define(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def header_row(ws, row, headers, widths=None):
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=text)
        style(c, f_head, HEAD_FILL, align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=BOX)
    ws.row_dimensions[row].height = 30
    if widths:
        for col, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w


def title(ws, text, sub):
    style(ws.cell(row=1, column=1, value=text), f_title)
    style(ws.cell(row=2, column=1, value=sub), f_sub)
    ws.row_dimensions[1].height = 28
    ws.sheet_view.showGridLines = False


# ---------- sheets ----------


def build_settings(wb):
    ws = wb.create_sheet("Settings")
    title(ws, "Settings & lists", "Edit the yellow cells. Lists feed the dropdowns on the other tabs.")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 4
    ws.column_dimensions["F"].width = 34
    ws.column_dimensions["G"].width = 4
    ws.column_dimensions["H"].width = 16
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 44

    style(ws["A4"], f_section).value = "Fee estimates"
    style(ws["A5"], f_label).value = "Typical buyer sales tax rate"
    style(ws["B5"], f_input, INPUT_FILL, "0.00%", border=BOX).value = engine()["taxRate"] / 100
    taxed = f"{tax_fee_names()} charge part of their fees on tax. " if tax_fee_names() else ""
    style(ws["A6"], f_note).value = f"{taxed}Typical US rates are 6-10%."
    define(wb, "TaxRate", "Settings!$B$5")

    def column_list(col, head, items, name, note):
        style(ws[f"{col}8"], f_section).value = head
        for i, item in enumerate(items):
            style(ws[f"{col}{9 + i}"], f_input, INPUT_FILL, border=BOX).value = item
        end = 9 + len(items) - 1
        define(wb, name, f"Settings!${col}$9:${col}${end}")
        style(ws[f"{col}{end + 1}"], f_note).value = note

    column_list("A", "Marketplaces", platform_names(), "Platforms",
                "Keep the order (fee formulas follow it). Renamed one? Find & Replace the old name on Inventory too.")
    column_list("D", "Sources", SOURCES, "Sources", "Where you buy. Rename to match your spots.")
    column_list("F", "Expense categories", EXPENSE_CATEGORIES, "ExpenseCategories", "Rename to suit your records.")

    style(ws["H8"], f_section).value = "Mileage rates"
    for col, text in zip("HIJ", ["Effective from", "Rate per mile", "Source"]):
        style(ws[f"{col}9"], f_head, HEAD_FILL, border=BOX).value = text
    # From src/data/mileage.mjs, the list the website's tracker page quotes.
    rates = engine()["mileage"]
    if len(rates) > 7:
        raise SystemExit("Keep at most 7 mileage rates (the sheet leaves room for 3 more): drop the oldest.")
    for i, rate in enumerate(rates):
        r = 10 + i
        style(ws[f"H{r}"], f_input, INPUT_FILL, DATE_FMT, border=BOX).value = date.fromisoformat(rate["from"])
        style(ws[f"I{r}"], f_input, INPUT_FILL, "$0.000", border=BOX).value = rate["rate"]
        style(ws[f"J{r}"], f_body, border=BOX).value = rate["label"]
    # Leave room for future rates: the named ranges cover 10 rows.
    for r in range(10 + len(rates), 20):
        style(ws[f"H{r}"], f_input, INPUT_FILL, DATE_FMT, border=BOX)
        style(ws[f"I{r}"], f_input, INPUT_FILL, "$0.000", border=BOX)
        style(ws[f"J{r}"], f_body, border=BOX)
    define(wb, "MileageStarts", "Settings!$H$10:$H$19")
    define(wb, "MileageRates", "Settings!$I$10:$I$19")
    source = engine()["mileageSource"].removeprefix("https://www.")
    style(ws["H20"], f_note).value = f"Add new IRS rates below the last row, oldest to newest. Source: {source}"
    return ws


def build_fees(wb):
    ws = wb.create_sheet("Fees")
    info = engine()
    title(ws, "Marketplace fees", f"US seller rates, verified {info['verified_label']}. Change a yellow cell and every estimate updates.")
    header_row(ws, 4, ["Marketplace", "Setting", "Value", "Name in formulas", "Notes"], [22, 36, 12, 22, 60])
    rates = info["rates"]
    names = {p["id"]: p["name"] for p in info["platforms"]}
    words = {**info["text"], "other": OTHER}
    for i, (name, platform, setting, source, fmt, note) in enumerate(FEES):
        r = 5 + i
        value = rates[platform][source] if platform else source
        style(ws[f"A{r}"], f_label, border=BOX).value = names[platform] if platform else OTHER
        style(ws[f"B{r}"], f_body, border=BOX).value = setting
        style(ws[f"C{r}"], f_input, INPUT_FILL, fmt, border=BOX).value = value
        style(ws[f"D{r}"], f_sub, border=BOX).value = name
        style(ws[f"E{r}"], f_body, border=BOX, align=Alignment(wrap_text=True, vertical="top")).value = note.format(**words)
        define(wb, name, f"Fees!$C${r}")
    r = 6 + len(FEES)
    style(ws[f"A{r}"], f_section).value = "Official fee pages"
    for i, p in enumerate(info["platforms"]):
        url = p["source"]
        style(ws[f"A{r + 1 + i}"], f_label).value = p["name"]
        c = style(ws[f"B{r + 1 + i}"], Font(name=FONT, size=10, color=BRAND, underline="single"))
        c.value = url
        c.hyperlink = url
    ws.freeze_panes = "A5"
    return ws


def fee_formula(r):
    """Estimated marketplace fee for inventory row r, one CHOOSE branch per Platforms entry."""
    p, s = f"$H{r}", f"$I{r}"
    total = f"$T{r}"  # order total incl. tax, from the helper column
    ebay = (
        f"ROUND(MIN({total},eBay_FVF_Cap)*eBay_FVF+MAX(0,{total}-eBay_FVF_Cap)*eBay_FVF_Over,2)"
        f"+IF({total}<=eBay_SmallOrderMax,eBay_OrderFeeSmall,eBay_OrderFee)"
    )
    posh = f"IF({p}<Posh_Threshold,Posh_Flat,ROUND({p}*Posh_Rate,2))"
    mercari = f"ROUND(({p}+{s})*Mercari_Rate,2)"
    depop = f"ROUND({total}*Depop_Proc,2)+Depop_ProcFixed"
    etsy = f"Etsy_Listing+ROUND(({p}+{s})*Etsy_Txn,2)+ROUND({total}*Etsy_Proc,2)+Etsy_ProcFixed"
    whatnot = f"ROUND({p}*Whatnot_Comm,2)+ROUND({total}*Whatnot_Proc,2)+Whatnot_ProcFixed"
    fb = f"MAX(FB_Min,ROUND(({p}+{s})*FB_Rate,2))"
    grailed = (
        f"IF({p}<Grailed_Threshold,MAX(Grailed_Min,ROUND({p}*Grailed_LowRate,2)),ROUND({p}*Grailed_Rate,2))"
        f"+ROUND(({p}+{s})*Grailed_Proc,2)+Grailed_ProcFixed"
    )
    tiktok = f"ROUND(({p}+{s})*TikTok_Rate,2)"
    other = f"ROUND(({p}+{s})*Other_Rate,2)+Other_Fixed"
    formulas = {"ebay": ebay, "poshmark": posh, "mercari": mercari, "depop": depop, "etsy": etsy,
                "whatnot": whatnot, "facebook": fb, "grailed": grailed, "tiktok": tiktok}
    ids = [pl["id"] for pl in engine()["platforms"]]
    if set(ids) != set(formulas):
        raise SystemExit(f"Fee formulas cover {sorted(formulas)} but the engine has {ids}: update fee_formula().")
    branches = ",".join([formulas[i] for i in ids] + [other])
    return f'=IF(OR($F{r}="",{p}=""),"",IFERROR(CHOOSE(MATCH($F{r},Platforms,0),{branches}),""))'


INV_HEADERS = [
    ("Item ID", 11, "input", None),
    ("Item", 30, "input", None),
    ("Source", 20, "input", None),
    ("Date bought", 12, "input", DATE_FMT),
    ("Cost", 10, "input", USD),
    ("Sold on", 20, "input", None),
    ("Date sold", 12, "input", DATE_FMT),
    ("Sale price", 11, "input", USD),
    ("Shipping charged", 11, "input", USD),
    ("Label cost", 10, "input", USD),
    ("Other costs", 10, "input", USD),
    ("Est. fees", 10, "calc", USD),
    ("Actual fees (optional)", 12, "input", USD),
    ("Fees", 10, "calc", USD),
    ("Payout", 11, "calc", USD),
    ("Profit", 11, "calc", USD),
    ("ROI", 9, "calc", PCT),
    ("Days to sell", 9, "calc", INT),
    ("Status", 10, "calc", None),
    ("Order total incl. est. tax", 12, "calc", USD),
]


def date_validation():
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True)
    dv.errorTitle = "Not a date"
    dv.error = f"Enter a date, for example {as_of(11):%m/%d/%Y}."
    return dv


def amount_validation(error="Enter an amount of 0 or more, without the $ sign."):
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True)
    dv.errorTitle = "Not a valid number"
    dv.error = error
    return dv


def build_inventory(wb, verify=False):
    ws = wb.create_sheet("Inventory")
    title(ws, "Inventory & sales", "One row per item. Fill the yellow cells; the gray cells calculate themselves. Row 5 is an example: type over it.")
    header_row(ws, 4, [h for h, *_ in INV_HEADERS], [w for _, w, *_ in INV_HEADERS])
    author = engine()["brand"]
    ws["L4"].comment = Comment("Estimated from the Fees tab. For exact records, type the fee from your payout in Actual fees.", author)
    ws["M4"].comment = Comment("Optional. When filled, it replaces the estimate for this sale.", author)
    ws["F4"].comment = Comment(
        "Pick from the list (Settings tab). Turns red when a sale has no marketplace from the list: "
        "its fees, payout and profit stay blank (even with Actual fees) until you pick one.", author)
    ignored = f" Ignored for {buyer_ships_names()}, where the buyer pays the label." if buyer_ships_names() else ""
    taxed = f" {tax_fee_names()} charge part of their fees on it." if tax_fee_names() else ""
    ws["I4"].comment = Comment(f"What the buyer paid you for shipping. Leave blank for free shipping.{ignored}", author)
    ws["J4"].comment = Comment(f"Postage you paid.{ignored}", author)
    ws["T4"].comment = Comment(f"Item + shipping + estimated buyer sales tax (Settings tab).{taxed}", author)

    for r in range(FIRST, LAST + 1):
        for col, (_, _, kind, fmt) in enumerate(INV_HEADERS, start=1):
            c = ws.cell(row=r, column=col)
            if kind == "input":
                style(c, f_input, INPUT_FILL, fmt, border=BOX)
            else:
                style(c, f_calc, CALC_FILL, fmt, border=BOX)
        ship = f"IF({buyer_ships_test(f'$F{r}')},0,N($I{r}))"
        label = f"IF({buyer_ships_test(f'$F{r}')},0,N($J{r}))"
        # Tax is rounded to the cent before fees apply, exactly like the website engine.
        ws[f"T{r}"] = f'=IF(OR($F{r}="",$H{r}=""),"",$H{r}+{ship}+ROUND(($H{r}+{ship})*TaxRate,2))'
        ws[f"L{r}"] = fee_formula(r)
        # No marketplace from the list (Sold on blank or unknown, so no estimate):
        # no fee either, even with Actual fees, so the sale stays out of every
        # total (the by-marketplace table couldn't place it). Sold on is highlighted.
        ws[f"N{r}"] = f'=IF(OR($H{r}="",$L{r}=""),"",IF($M{r}<>"",$M{r},$L{r}))'
        ws[f"O{r}"] = f'=IF(OR($H{r}="",$N{r}=""),"",$H{r}+{ship}-$N{r})'
        ws[f"P{r}"] = f'=IF($O{r}="","",$O{r}-{label}-N($E{r})-N($K{r}))'
        ws[f"Q{r}"] = f'=IF(OR($P{r}="",N($E{r})=0),"",$P{r}/$E{r})'
        ws[f"R{r}"] = f'=IF(OR($D{r}="",$G{r}=""),"",$G{r}-$D{r})'
        # Half-entered sales are flagged and left out of the totals: a price without
        # a date sold can't be put in a year, a date without a price has no amount.
        ws[f"S{r}"] = (
            f'=IF($G{r}<>"",IF($H{r}="","Needs price","Sold"),'
            f'IF($H{r}<>"","Needs date",IF(OR($A{r}<>"",$B{r}<>"",$E{r}<>""),"In stock","")))'
        )

    example = ["TV-0001", "Patagonia Better Sweater, men's M", "Goodwill", as_of(23), 8, "Poshmark", as_of(11), 45, None, None, 0.5, None, None]
    for col, value in enumerate(example, start=1):
        if value is not None:
            ws.cell(row=FIRST, column=col, value=value)

    if verify:
        # Every marketplace at every VERIFY_PRICES x VERIFY_SHIPPING point, for the cross-check.
        for i, (market, price, ship) in enumerate(verify_cases()):
            r = FIRST + 1 + i
            row = ["VERIFY", market, "Other", as_of(24), VERIFY_COST, market, as_of(5), price, ship, VERIFY_LABEL]
            for col, value in enumerate(row, start=1):
                ws.cell(row=r, column=col, value=value)
        # Sales whose fee can't be worked out: payout and profit must stay blank.
        # With Actual fees too: without a listed marketplace they don't count either.
        for i, market in enumerate(NO_FEE_CASES):
            r = no_fee_row(i)
            row = ["VERIFY", "no fee", "Other", as_of(24), VERIFY_COST, market, as_of(5), 45, 5, VERIFY_LABEL, None, None, 3]
            for col, value in enumerate(row, start=1):
                if value is not None:  # L is a formula
                    ws.cell(row=r, column=col, value=value)
        for i, market in enumerate(NO_DATE_CASES):
            row = ["VERIFY", "no date", "Other", as_of(24), VERIFY_COST, market, None, 45, 5, VERIFY_LABEL]
            for col, value in enumerate(row, start=1):
                ws.cell(row=no_date_row(i), column=col, value=value)
        for i, market in enumerate(NO_PRICE_CASES):
            row = ["VERIFY", "no price", "Other", as_of(24), VERIFY_COST, market, as_of(5), None, 5, VERIFY_LABEL]
            for col, value in enumerate(row, start=1):
                ws.cell(row=no_price_row(i), column=col, value=value)

    # showErrorMessage makes Excel and Sheets refuse bad entries (openpyxl's default is off).
    dv_platform = DataValidation(type="list", formula1="=Platforms", allow_blank=True, showErrorMessage=True)
    dv_platform.error = "Pick a marketplace from the list (edit the list on the Settings tab)."
    dv_platform.errorTitle = "Unknown marketplace"
    dv_source = DataValidation(type="list", formula1="=Sources", allow_blank=True, showErrorMessage=False)
    dv_date = date_validation()
    dv_money = amount_validation()
    for dv in (dv_platform, dv_source, dv_date, dv_money):
        ws.add_data_validation(dv)
    dv_platform.add(f"F{FIRST}:F{LAST}")
    dv_source.add(f"C{FIRST}:C{LAST}")
    dv_date.add(f"D{FIRST}:D{LAST}")
    dv_date.add(f"G{FIRST}:G{LAST}")
    for col in "EHIJKM":
        dv_money.add(f"{col}{FIRST}:{col}{LAST}")

    red = Font(name=FONT, size=10, color="B3261E", bold=True)
    ws.conditional_formatting.add(
        f"F{FIRST}:F{LAST}",
        FormulaRule(formula=[f'AND($H{FIRST}<>"",$N{FIRST}="")'], font=red, fill=PatternFill("solid", fgColor="FBE3E1")),
    )
    ws.conditional_formatting.add(
        f"G{FIRST}:G{LAST}",
        FormulaRule(formula=[f'AND($H{FIRST}<>"",$G{FIRST}="")'], font=red, fill=PatternFill("solid", fgColor="FBE3E1")),
    )
    ws.conditional_formatting.add(
        f"H{FIRST}:H{LAST}",
        FormulaRule(formula=[f'AND($G{FIRST}<>"",$H{FIRST}="")'], font=red, fill=PatternFill("solid", fgColor="FBE3E1")),
    )
    ws.conditional_formatting.add(f"P{FIRST}:P{LAST}", CellIsRule(operator="lessThan", formula=["0"], font=red))
    ws.conditional_formatting.add(
        f"S{FIRST}:S{LAST}", FormulaRule(formula=[f'$S{FIRST}="Sold"'], font=Font(name=FONT, size=10, color=BRAND, bold=True))
    )
    ws.conditional_formatting.add(
        f"I{FIRST}:J{LAST}",
        FormulaRule(formula=[buyer_ships_test(f"$F{FIRST}")], font=Font(name=FONT, size=10, color="9AA5A0", strike=True)),
    )
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:T{LAST}"
    return ws


def build_expenses(wb):
    ws = wb.create_sheet("Expenses")
    title(ws, "Business expenses", "Costs that are not tied to one item. Row 5 is an example: type over it. Totals by category are on the Dashboard.")
    header_row(ws, 4, ["Date", "Category", "Description", "Amount", "Notes", "Tax year"], [12, 30, 34, 12, 30, 10])
    for r in range(FIRST, LAST + 1):
        for col, fmt in zip("ABCDE", [DATE_FMT, None, None, USD, None]):
            style(ws[f"{col}{r}"], f_input, INPUT_FILL, fmt, border=BOX)
        # For filtering by year. Every row also needs a formula for LibreOffice's
        # re-save (recalculate) to keep all the prebuilt rows and their checks.
        style(ws[f"F{r}"], f_calc, CALC_FILL, "0", border=BOX).value = f'=IF($A{r}="","",YEAR($A{r}))'
    for col, value in zip("ABCDE", [as_of(22), "Shipping supplies", "Poly mailers, 100 pack", 14.99, "Receipt in email"]):
        ws[f"{col}{FIRST}"] = value

    dv_cat = DataValidation(type="list", formula1="=ExpenseCategories", allow_blank=True, showErrorMessage=False)
    dv_date, dv_amount = date_validation(), amount_validation()
    for dv in (dv_cat, dv_date, dv_amount):
        ws.add_data_validation(dv)
    dv_cat.add(f"B{FIRST}:B{LAST}")
    dv_date.add(f"A{FIRST}:A{LAST}")
    dv_amount.add(f"D{FIRST}:D{LAST}")
    ws.auto_filter.ref = f"A4:F{LAST}"  # e.g. filter Tax year at tax time

    ws.freeze_panes = "A5"
    return ws


def mileage_verify_cases():
    """Trips on each IRS rate's first day, the day before, and well after the last."""
    starts = [date.fromisoformat(r["from"]) for r in engine()["mileage"]]
    days = sorted({d for s in starts for d in (s - timedelta(days=1), s)} | {starts[-1] + timedelta(days=200)})
    return [(d, 12.5) for d in days]


def build_mileage(wb, verify=False):
    ws = wb.create_sheet("Mileage")
    title(ws, "Business mileage", "Sourcing runs, post office trips, supply runs. The IRS rate is picked by date. Row 5 is an example: type over it.")
    header_row(ws, 4, ["Date", "Purpose", "From", "To", "Miles", "Rate", "Deduction"], [12, 30, 20, 20, 9, 9, 12])
    for r in range(FIRST, LAST + 1):
        for col, fmt in zip("ABCDE", [DATE_FMT, None, None, None, "#,##0.0"]):
            style(ws[f"{col}{r}"], f_input, INPUT_FILL, fmt, border=BOX)
        style(ws[f"F{r}"], f_calc, CALC_FILL, "$0.000", border=BOX).value = (
            f'=IF(OR($A{r}="",$E{r}=""),"",IFERROR(INDEX(MileageRates,MATCH($A{r},MileageStarts,1)),""))'
        )
        style(ws[f"G{r}"], f_calc, CALC_FILL, USD, border=BOX).value = f'=IF($F{r}="","",ROUND($E{r}*$F{r},2))'
    for col, value in zip("ABCDE", [as_of(23), "Sourcing: thrift route", "Home", "Goodwill, Main St", 18.4]):
        ws[f"{col}{FIRST}"] = value
    if verify:
        for i, (day, miles) in enumerate(mileage_verify_cases()):
            for col, value in zip("ABE", [day, "VERIFY", miles]):
                ws[f"{col}{FIRST + 1 + i}"] = value
    dv_date, dv_miles = date_validation(), amount_validation("Enter the miles as a number, 0 or more.")
    for dv in (dv_date, dv_miles):
        ws.add_data_validation(dv)
    dv_date.add(f"A{FIRST}:A{LAST}")
    dv_miles.add(f"E{FIRST}:E{LAST}")
    ws.freeze_panes = "A5"
    return ws


INV = "Inventory!"


def inv(col):
    """The whole data range of one Inventory column, for SUMIFS/COUNTIFS."""
    return f"{INV}${col}${FIRST}:${col}${RANGE_END}"


def between(dates, start, end):
    """COUNTIFS/SUMIFS criteria pair: a date in `dates` is on or after `start` and before `end`."""
    return f'{dates},">="&{start},{dates},"<"&{end}'


def in_dash_year(dates):
    """COUNTIFS/SUMIFS criteria pair: a date in `dates` falls in the Dashboard year."""
    return between(dates, "DATE(DashYear,1,1)", "DATE(DashYear+1,1,1)")


SOLD_IN_YEAR = in_dash_year(inv("G"))


def count_text(count, one, many):
    """Formula text: '1 sale is' / 'N sales are' for the count in `count`."""
    return f'IF({count}=1,"1 {one}",{count}&" {many}")'


def build_checks(wb):
    """Settings > Data checks: sales left out of the Dashboard totals, counted
    once in two labelled cells that the Dashboard warning reads."""
    ws = wb["Settings"]
    ws.column_dimensions["L"].width = 64
    ws.column_dimensions["M"].width = 10
    style(ws["L8"], f_section).value = "Data checks"
    checks = [
        ("NoMarketplaceSales", "Sales in the Dashboard year with no marketplace from the Settings list",
         f'=COUNTIFS({inv("H")},"<>",{inv("N")},"",{SOLD_IN_YEAR})'),
        ("HalfEnteredSales", "Sales missing a date sold or a price (any year)",
         f'=COUNTIFS({inv("H")},"<>",{inv("G")},"")'
         f'+COUNTIFS({inv("G")},"<>",{inv("H")},"")'),
    ]
    for i, (name, label, formula) in enumerate(checks):
        style(ws[f"L{9 + i}"], f_body, border=BOX).value = label
        style(ws[f"M{9 + i}"], f_calc, CALC_FILL, "0", border=BOX).value = formula  # "0", not "-": the note says both should be 0
        define(wb, name, f"Settings!$M${9 + i}")
    style(ws["L11"], f_note).value = "Both should be 0. The Dashboard warns when they are not."


def warning_formula(no_marketplace, half_entered):
    """Dashboard E3: the sales left out of the totals, in words (blank when there are none)."""
    n, h = no_marketplace, half_entered
    return (
        f'=IF({n}+{h}=0,"",TRIM('
        f'IF({n}>0,{count_text(n, "sale", "sales")}&" in "&DashYear&IF({n}=1," has"," have")&" no marketplace from the Settings list. ","")'
        f'&IF({h}>0,{count_text(h, "sale is", "sales are")}&" missing a date sold or a price. ","")'
        '&"Left out of the totals: see the red cells on Inventory."))'
    )


# Counts (no marketplace, half-entered) the verify build writes the warning for,
# so singular, plural and blank wording are all checked.
WARNING_CASES = [(0, 0), (1, 0), (0, 1), (1, 1), (2, 3)]


def build_warning_checks(wb):
    """Verify build only: the Dashboard warning worded for each of WARNING_CASES."""
    ws = wb.create_sheet("Warning checks")
    for i, (n, h) in enumerate(WARNING_CASES, start=1):
        ws[f"A{i}"], ws[f"B{i}"] = n, h
        ws[f"C{i}"] = warning_formula(f"A{i}", f"B{i}")


def build_dashboard(wb, verify=False):
    ws = wb.create_sheet("Dashboard", 1)
    title(ws, "Dashboard", "Everything below updates from your Inventory, Expenses and Mileage tabs.")
    for col, w in zip("ABCDEFGH", [26, 18, 18, 18, 14, 14, 14, 14]):
        ws.column_dimensions[col].width = w

    style(ws["A3"], f_label).value = "Year"
    # The current year when the file is opened; type a year over it to see another.
    # The verify build pins the year its sample sales are in.
    year = as_of(5).year if verify else "=YEAR(TODAY())"
    style(ws["B3"], Font(name=FONT, size=12, bold=True, color="0000FF"), INPUT_FILL, "0", border=BOX).value = year
    style(ws["C3"], f_note).value = "This year. Type a year over it (last year's, at tax time)."
    define(wb, "DashYear", "Dashboard!$B$3")

    # Totals count only sales whose fee could be worked out (a number in Fees);
    # the others are left out entirely (sales, costs and counts) and flagged in E3.
    counted = f'{inv("N")},">=0"'
    in_year = f"{SOLD_IN_YEAR},{counted}"

    # KPI tiles
    # Sales left out of the totals (counted on Settings > Data checks).
    style(ws["E3"], Font(name=FONT, size=10, bold=True, color="B3261E")).value = warning_formula("NoMarketplaceSales", "HalfEnteredSales")
    kpis = [
        ("Items sold", f"=COUNTIFS({in_year})", INT),
        ("Sales incl. shipping", f"=SUMIFS({inv('O')},{in_year})+SUMIFS({inv('N')},{in_year})", USD0),
        ("Marketplace fees", f"=SUMIFS({inv('N')},{in_year})", USD0),
        ("Profit", f"=SUMIFS({inv('P')},{in_year})", USD0),
        ("Avg profit per sale", '=IF(A6=0,"",D6/A6)', USD),
        ("Avg days to sell", f'=IFERROR(AVERAGEIFS({inv("R")},{in_year}),"")', "0"),
        ("Sell-through (all time)", f'=IFERROR(COUNTIF({inv("S")},"Sold")/(COUNTIF({inv("S")},"Sold")+COUNTIF({inv("S")},"In stock")),"")', PCT),
        ("Unsold stock at cost", f'=SUMIFS({inv("E")},{inv("S")},"In stock")', USD0),
    ]
    # Two rows of four tiles: labels on 5/8, values on 6/9.
    for i, (label, formula, fmt) in enumerate(kpis):
        row = 5 if i < 4 else 8
        col = "ABCD"[i % 4]
        lab = ws[f"{col}{row}"]
        val = ws[f"{col}{row + 1}"]
        style(lab, f_sub, SOFT_FILL, align=Alignment(wrap_text=True, vertical="top"))
        lab.value = label
        style(val, f_kpi, SOFT_FILL, fmt, align=Alignment(horizontal="left"))
        val.value = formula
    for r in (5, 8):
        ws.row_dimensions[r].height = 18
    for r in (6, 9):
        ws.row_dimensions[r].height = 26

    # Monthly table
    style(ws["A12"], f_section).value = "By month"
    for col, text in zip("ABCDEF", ["Month", "Items sold", "Sales", "Fees", "Profit", "Margin"]):
        style(ws[f"{col}13"], f_head, HEAD_FILL, border=BOX, align=Alignment(horizontal="center")).value = text
    for m in range(1, 13):
        r = 13 + m
        month = between(inv("G"), f"$A{r}", f"DATE(YEAR($A{r}),MONTH($A{r})+1,1)") + f",{counted}"
        style(ws[f"A{r}"], f_label, fmt="mmmm", border=BOX, align=Alignment(horizontal="left")).value = f"=DATE(DashYear,{m},1)"
        style(ws[f"B{r}"], f_calc, fmt=INT, border=BOX).value = f"=COUNTIFS({month})"
        style(ws[f"C{r}"], f_calc, fmt=USD0, border=BOX).value = f"=SUMIFS({inv('O')},{month})+SUMIFS({inv('N')},{month})"
        style(ws[f"D{r}"], f_calc, fmt=USD0, border=BOX).value = f"=SUMIFS({inv('N')},{month})"
        style(ws[f"E{r}"], f_calc, fmt=USD0, border=BOX).value = f"=SUMIFS({inv('P')},{month})"
        style(ws[f"F{r}"], f_calc, fmt=PCT, border=BOX).value = f'=IF(C{r}=0,"",E{r}/C{r})'
    style(ws["A26"], f_label, SOFT_FILL, border=BOX).value = "Year total"
    for col, fmt in zip("BCDE", [INT, USD0, USD0, USD0]):
        style(ws[f"{col}26"], f_label, SOFT_FILL, fmt, border=BOX).value = f"=SUM({col}14:{col}25)"
    style(ws["F26"], f_label, SOFT_FILL, PCT, border=BOX).value = '=IF(C26=0,"",E26/C26)'

    chart = BarChart()
    chart.type = "col"
    chart.title = "Profit by month"
    chart.style = 10
    chart.legend = None
    chart.y_axis.numFmt = "$#,##0"
    chart.y_axis.majorGridlines = None
    chart.x_axis.number_format = "mmm"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    data = Reference(ws, min_col=5, min_row=13, max_row=25)
    cats = Reference(ws, min_col=1, min_row=14, max_row=25)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.series[0].graphicalProperties.solidFill = BRAND
    chart.height = 7.5
    chart.width = 15
    ws.add_chart(chart, "H12")

    # By marketplace
    style(ws["A29"], f_section).value = "By marketplace (Dashboard year)"
    heads = ["Marketplace", "Items sold", "Sales", "Fees", "Profit", "Avg profit", "Avg days to sell"]
    for col, text in zip("ABCDEFG", heads):
        style(ws[f"{col}30"], f_head, HEAD_FILL, border=BOX, align=Alignment(horizontal="center", wrap_text=True)).value = text
    for i in range(len(platform_names())):
        r = 31 + i
        by = f"{inv('F')},$A{r},{in_year}"
        style(ws[f"A{r}"], f_label, border=BOX).value = f"=INDEX(Platforms,{i + 1})"
        style(ws[f"B{r}"], f_calc, fmt=INT, border=BOX).value = f"=COUNTIFS({by})"
        style(ws[f"C{r}"], f_calc, fmt=USD0, border=BOX).value = f"=SUMIFS({inv('O')},{by})+SUMIFS({inv('N')},{by})"
        style(ws[f"D{r}"], f_calc, fmt=USD0, border=BOX).value = f"=SUMIFS({inv('N')},{by})"
        style(ws[f"E{r}"], f_calc, fmt=USD0, border=BOX).value = f"=SUMIFS({inv('P')},{by})"
        style(ws[f"F{r}"], f_calc, fmt=USD, border=BOX).value = f'=IF(B{r}=0,"",E{r}/B{r})'
        style(ws[f"G{r}"], f_calc, fmt="0", border=BOX).value = f'=IFERROR(AVERAGEIFS({inv("R")},{by}),"")'
    last_market = 30 + len(platform_names())

    # By source (all time)
    top = last_market + 3
    style(ws[f"A{top}"], f_section).value = "By source (all time)"
    heads = ["Source", "Items bought", "Items sold", "Cost", "Profit", "Sell-through"]
    for col, text in zip("ABCDEF", heads):
        style(ws[f"{col}{top + 1}"], f_head, HEAD_FILL, border=BOX, align=Alignment(horizontal="center")).value = text
    for i in range(len(SOURCES)):
        r = top + 2 + i
        style(ws[f"A{r}"], f_label, border=BOX).value = f"=INDEX(Sources,{i + 1})"
        style(ws[f"B{r}"], f_calc, fmt=INT, border=BOX).value = f"=COUNTIFS({inv('C')},$A{r})"
        style(ws[f"C{r}"], f_calc, fmt=INT, border=BOX).value = f'=COUNTIFS({inv("C")},$A{r},{inv("S")},"Sold")'
        style(ws[f"D{r}"], f_calc, fmt=USD0, border=BOX).value = f"=SUMIFS({inv('E')},{inv('C')},$A{r})"
        style(ws[f"E{r}"], f_calc, fmt=USD0, border=BOX).value = (
            f'=SUMIFS({inv("P")},{inv("C")},$A{r},{inv("S")},"Sold")'
        )
        style(ws[f"F{r}"], f_calc, fmt=PCT, border=BOX).value = f'=IF(B{r}=0,"",C{r}/B{r})'
    last_source = top + 1 + len(SOURCES)

    # Expenses by category
    e = last_source + 3
    style(ws[f"A{e}"], f_section).value = "Expenses by category (Dashboard year)"
    for col, text in zip("AB", ["Category", "Total"]):
        style(ws[f"{col}{e + 1}"], f_head, HEAD_FILL, border=BOX, align=Alignment(horizontal="center")).value = text
    exp = in_dash_year(f"Expenses!$A${FIRST}:$A${RANGE_END}")
    for i in range(len(EXPENSE_CATEGORIES)):
        r = e + 2 + i
        style(ws[f"A{r}"], f_label, border=BOX).value = f"=INDEX(ExpenseCategories,{i + 1})"
        style(ws[f"B{r}"], f_calc, fmt=USD, border=BOX).value = (
            f"=SUMIFS(Expenses!$D${FIRST}:$D${RANGE_END},Expenses!$B${FIRST}:$B${RANGE_END},$A{r},{exp})"
        )
    exp_total = e + 2 + len(EXPENSE_CATEGORIES)
    style(ws[f"A{exp_total}"], f_label, SOFT_FILL, border=BOX).value = "All expenses"
    style(ws[f"B{exp_total}"], f_label, SOFT_FILL, USD, border=BOX).value = f"=SUMIFS(Expenses!$D${FIRST}:$D${RANGE_END},{exp})"

    # Tax-time summary
    t = exp_total + 3
    style(ws[f"A{t}"], f_section).value = "Tax-time summary (Dashboard year)"
    rows = [
        ("Gross sales (item + shipping charged)", "=C26"),
        ("Marketplace fees", "=-D26"),
        ("Cost of items sold", f"=-SUMIFS({inv('E')},{in_year})"),
        # Labels on marketplaces where the buyer pays them are ignored, as on each row.
        ("Shipping labels", f"=-(SUMIFS({inv('J')},{in_year})" + "".join(
            f"-SUMIFS({inv('J')},{inv('F')},INDEX(Platforms,{i}),{in_year})" for i in buyer_ships()
        ) + ")"),
        ("Other per-item costs", f"=-SUMIFS({inv('K')},{in_year})"),
        ("Business expenses (Expenses tab)", f"=-B{exp_total}"),
        ("Mileage deduction (Mileage tab)", f'=-SUMIFS(Mileage!$G${FIRST}:$G${RANGE_END},{in_dash_year(f"Mileage!$A${FIRST}:$A${RANGE_END}")})'),
    ]
    for i, (label, formula) in enumerate(rows):
        r = t + 1 + i
        style(ws[f"A{r}"], f_body, border=BOX).value = label
        style(ws[f"B{r}"], f_calc, fmt=USD, border=BOX).value = formula
    r = t + 1 + len(rows)
    style(ws[f"A{r}"], f_label, SOFT_FILL, border=BOX).value = "Estimated net profit"
    style(ws[f"B{r}"], f_label, SOFT_FILL, USD, border=BOX).value = f"=SUM(B{t + 1}:B{r - 1})"
    style(ws[f"A{r + 1}"], f_note).value = "An estimate to organize your records, not tax advice. Confirm with a tax professional or your tax software."
    return ws


def build_start(wb):
    ws = wb.active
    ws.title = "Start Here"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 96
    style(ws["B2"], f_title).value = engine()["product"]
    style(ws["B3"], f_sub).value = f"Track every flip from haul to payout. Marketplace fees verified {engine()['verified_label']}."

    style(ws["B5"], f_section).value = "Get started in five minutes"
    steps = [
        "Inventory tab: add a row when you buy something. Item ID, description, source, date bought and cost.",
        "When it sells, fill in Sold on, Date sold, Sale price, Shipping charged and Label cost. Fees, payout, profit, ROI and days to sell appear on their own.",
        "Optional: type the exact fee from your payout into Actual fees. The tracker uses it instead of the estimate.",
        "Log business costs on the Expenses tab and business driving on the Mileage tab.",
        "Open the Dashboard for this year's monthly profit, best marketplaces and sources, and tax-time totals. Type another year at the top to see it.",
    ]
    for i, text in enumerate(steps):
        c = style(ws[f"B{6 + i}"], f_body, align=Alignment(wrap_text=True, vertical="top"))
        c.value = f"{i + 1}.  {text}"
        ws.row_dimensions[6 + i].height = 30

    style(ws["B12"], f_section).value = "Cell colors"
    style(ws["B13"], f_input, INPUT_FILL, border=BOX).value = "Yellow with blue text: type here"
    style(ws["B14"], f_calc, CALC_FILL, border=BOX).value = "Gray: calculated for you. Leave these alone."

    style(ws["B16"], f_section).value = "Good to know"
    tips = [
        *([f"{buyer_ships_names()}: the buyer pays for the label, so Shipping charged and Label cost are ignored (shown struck through) on those rows."]
          if buyer_ships_names() else []),
        f"Local cash sale? Pick {OTHER} as the marketplace; its fee is 0 unless you change it on the Fees tab.",
        "A red cell on Inventory means a half-entered sale (no marketplace, date sold or price). It stays out of the Dashboard totals until you fill it in; the Dashboard says how many there are.",
        "Row 5 on each log is an example. Type over it, or right-click its row number and choose Delete to remove the whole row. Don't clear the gray cells on their own: they hold that row's formulas (copy one down from the row above to repair).",
        f"There are {DATA_ROWS:,} ready rows on each log. Need more? Copy the last row and paste it below; the formulas come along, and the Dashboard counts rows up to {RANGE_END:,}.",
        "Fees change: update the rates on the Fees tab. Sales with Actual fees filled in keep their exact numbers.",
        "Rename marketplaces, sources and expense categories on the Settings tab. Keep the marketplace order, because the fee formulas follow it.",
        "Google Sheets: upload this file to Google Drive, then open it with Google Sheets. Everything works except the chart styling.",
    ]
    for i, text in enumerate(tips):
        c = style(ws[f"B{17 + i}"], f_body, align=Alignment(wrap_text=True, vertical="top"))
        c.value = f"•  {text}"
        ws.row_dimensions[17 + i].height = 28

    note = 17 + len(tips) + 1  # below the tips, however many there are
    style(ws[f"B{note}"], f_note, align=Alignment(wrap_text=True)).value = (
        "Fee figures are estimates based on each marketplace's published US fees; promotions, store plans and category "
        "exceptions can change them. Tax totals organize your records and are not tax advice. "
        "Licensed for use in your own reselling business. Please do not share or resell this file."
    )
    ws.row_dimensions[note].height = 40
    return ws


def build(verify=False):
    wb = Workbook()
    build_start(wb)
    build_inventory(wb, verify)
    build_expenses(wb)
    build_mileage(wb, verify)
    build_fees(wb)
    build_settings(wb)
    build_checks(wb)
    build_dashboard(wb, verify)
    if verify:
        build_warning_checks(wb)
    wb.calculation.fullCalcOnLoad = True
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = BRAND if ws.title in ("Start Here", "Dashboard") else "B8C4BF"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^A-Za-z0-9]+", "-", engine()["product"]).strip("-")  # "ThreadVet Reseller Tracker" -> ThreadVet-Reseller-Tracker
    name = f"{stem}-verify.xlsx" if verify else f"{stem}.xlsx"
    path = OUT_DIR / name
    wb.save(path)
    return path


def recalculate(path):
    """Let LibreOffice compute every formula and store the results in place."""
    if not shutil.which("soffice"):
        print("soffice not found: skipping recalculation (values compute when the file is opened)")
        return path
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            ["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/profile",
             "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", tmp, str(path)],
            check=True, capture_output=True, timeout=300,
        )
        shutil.move(Path(tmp) / path.name, path)
    return path


if __name__ == "__main__":
    print(recalculate(build()))
