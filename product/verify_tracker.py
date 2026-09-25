"""
End-to-end check that the spreadsheet's fee formulas agree with the website.

Builds a copy of the tracker with a grid of sales on every marketplace,
lets LibreOffice calculate it, then compares every estimated fee with the
fee engine in src/engine (via Node). Also checks the workbook has no formula
errors. Requires LibreOffice Calc (`soffice`) and Node.

    python3 product/verify_tracker.py
"""

import json
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
import build_tracker  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
ENGINE = """
import { normalizeInputs, evaluate, DEFAULTS } from './src/engine/calc.mjs';
import { PLATFORMS } from './src/engine/fees.mjs';
const byName = Object.fromEntries(PLATFORMS.map((p) => [p.name, p]));
const [cases, cost, label] = JSON.parse(process.argv[1]);
console.log(JSON.stringify(cases.map(([name, price, ship]) => {
  // "Other / Local" has no engine platform: no fees, and every input counts.
  if (!byName[name]) return { fees: 0, payout: Math.round((price + ship) * 100), profit: Math.round((price + ship - label - cost) * 100) };
  // The sheet's default tax rate (Settings) is the calculator's default too.
  const r = evaluate(byName[name], normalizeInputs({ price, ship, label, cost, taxRate: DEFAULTS.taxRate }));
  return { fees: r.feeTotal, payout: r.payout, profit: r.profit };
})));
"""
ERRORS = ("#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")


def check_features(path):
    """Problems with spreadsheet features that must survive the LibreOffice re-save."""
    wb = load_workbook(path)
    problems = []
    expected_sheets = ["Start Here", "Dashboard", "Inventory", "Expenses", "Mileage", "Fees", "Settings"]
    if wb.sheetnames != expected_sheets:
        problems.append(f"sheets are {wb.sheetnames}")
    names = set(wb.defined_names.keys())
    for name in ["Platforms", "Sources", "ExpenseCategories", "TaxRate", "DashYear", "MileageRates", "eBay_FVF", "Grailed_Rate"]:
        if name not in names:
            problems.append(f"named range {name} missing")
    inv = wb["Inventory"]
    lists = {str(dv.formula1).lstrip("="): str(dv.sqref) for dv in inv.data_validations.dataValidation if dv.type == "list"}
    if not lists.get("Platforms", "").startswith("F5") or not lists.get("Sources", "").startswith("C5"):
        problems.append(f"Inventory dropdowns changed: {lists}")
    if not any(dv.type == "date" for dv in inv.data_validations.dataValidation):
        problems.append("Inventory date validation missing")
    if len(list(inv.conditional_formatting)) < 3:
        problems.append("Inventory conditional formatting missing")
    if not [c for row in inv.iter_rows(min_row=4, max_row=4) for c in row if c.comment]:
        problems.append("Inventory header comments missing")
    if inv.freeze_panes != "C5":
        problems.append(f"Inventory freeze panes are {inv.freeze_panes}")
    if not wb["Dashboard"]._charts:
        problems.append("Dashboard chart missing")
    return problems


def main():
    if not shutil.which("soffice"):
        sys.exit("LibreOffice Calc (soffice) is required to calculate the workbook. Install it and re-run.")
    src = build_tracker.recalculate(build_tracker.build(verify=True))
    wb = load_workbook(src, data_only=True)
    src.unlink()

    problems = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERRORS:
                    problems.append(f"{ws.title}!{c.coordinate} = {c.value}")

    cases = build_tracker.verify_cases()
    payload = [cases, build_tracker.VERIFY_COST, build_tracker.VERIFY_LABEL]
    expected = json.loads(subprocess.run(
        ["node", "--input-type=module", "-e", ENGINE, json.dumps(payload)],
        cwd=ROOT, check=True, capture_output=True, text=True,
    ).stdout)
    inv = wb["Inventory"]
    columns = {"fees": "L", "payout": "O", "profit": "P"}
    for i, ((name, price, ship), want) in enumerate(zip(cases, expected)):
        r = build_tracker.FIRST + 1 + i
        for key, col in columns.items():
            value = inv[f"{col}{r}"].value
            if not isinstance(value, (int, float)):
                problems.append(f"{name} at ${price} + ${ship} shipping, {key}: sheet shows {value!r}")
                continue
            got = round(value * 100)
            if got != want[key]:
                problems.append(f"{name} at ${price} + ${ship} shipping, {key}: sheet {got} vs engine {want[key]} cents")
    # A sale with no usable marketplace gets no payout or profit, and the Dashboard says so.
    before = len(problems)
    for i, market in enumerate(build_tracker.NO_FEE_CASES):
        r = build_tracker.no_fee_row(i)
        for col in "LOP":
            if inv[f"{col}{r}"].value not in (None, ""):
                problems.append(f"Sold on {market!r}: {col}{r} should be blank, sheet shows {inv[f'{col}{r}'].value!r}")
    warning = wb["Dashboard"]["E3"].value or ""
    if not warning.startswith(f"{len(build_tracker.NO_FEE_CASES)} sale(s) left out"):
        problems.append(f"Dashboard warning for sales without a marketplace reads {warning!r}")
    # ...and they are left out everywhere: counts and every cost line of the tax summary.
    dash = wb["Dashboard"]
    counted = [r for r in range(build_tracker.FIRST, build_tracker.LAST + 1) if isinstance(inv[f"N{r}"].value, (int, float))]
    label_of = {dash[f"A{r}"].value: dash[f"B{r}"].value for r in range(1, dash.max_row + 1)}
    expected_totals = {
        "Items sold": (dash["A6"].value, len(counted)),
        "Cost of items sold": (label_of.get("Cost of items sold"), -sum(inv[f"E{r}"].value or 0 for r in counted)),
        "Other per-item costs": (label_of.get("Other per-item costs"), -sum(inv[f"K{r}"].value or 0 for r in counted)),
    }
    for name, (got, want) in expected_totals.items():
        if not isinstance(got, (int, float)) or abs(got - want) > 0.005:
            problems.append(f"Dashboard {name}: sheet {got!r}, expected {want!r} (sales with a fee only)")
    no_fee_ok = len(problems) == before

    # Mileage: the rate each trip gets must be the IRS rate in force that day (src/data/mileage.mjs).
    rates = [(build_tracker.date.fromisoformat(r["from"]), r["rate"]) for r in build_tracker.engine()["mileage"]]
    trips = build_tracker.mileage_verify_cases()
    miles_ws = wb["Mileage"]
    for i, (day, miles) in enumerate(trips):
        r = build_tracker.FIRST + 1 + i
        want = next((rate for start, rate in reversed(rates) if start <= day), None)
        got_rate, got_deduction = miles_ws[f"F{r}"].value, miles_ws[f"G{r}"].value
        if want is None:
            if got_rate not in (None, ""):
                problems.append(f"Mileage on {day}: before the first IRS rate, sheet shows {got_rate!r}")
        elif got_rate != want or round(got_deduction * 100) != round(miles * want * 100):
            problems.append(f"Mileage on {day}: sheet {got_rate!r}/{got_deduction!r} vs IRS {want} x {miles} miles")

    print(f"Compared fees, payout and profit on {len(cases)} sales ({len(build_tracker.platform_names())} marketplaces x "
          f"{len(build_tracker.VERIFY_PRICES)} prices x {len(build_tracker.VERIFY_SHIPPING)} shipping amounts).")

    print(f"Checked the IRS mileage rate picked on {len(trips)} trip dates around each rate change.")
    if no_fee_ok:
        print("Sales without a usable marketplace stay out of payout, profit and the totals, with a Dashboard warning.")

    product = build_tracker.recalculate(build_tracker.build())
    feature_problems = check_features(product)
    problems += feature_problems
    if not feature_problems:
        print(f"Delivered file keeps its dropdowns, named ranges, formatting, comments and chart: {product.name}")

    if problems:
        print("\n".join(["", "FAILED:", *problems]))
        sys.exit(1)
    print("Every figure matches the website engine to the cent; no formula errors.")


if __name__ == "__main__":
    main()
